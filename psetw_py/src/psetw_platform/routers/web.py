"""Server-rendered PS-ETW web UI routes."""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from calendar import SUNDAY, Calendar
from collections.abc import Iterable
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Annotated
from urllib.parse import urlencode

from fastapi import APIRouter, File, Form, Request, Response, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from psetw_platform.core.security import create_access_token, decode_access_token, verify_password
from psetw_platform.dependencies import DBSession
from psetw_platform.models import (
    ActionItem,
    ActionPriority,
    ActionStatus,
    Activity,
    ActivityStatus,
    Baseline,
    EngineDocument,
    EodLog,
    Project,
    User,
    UserRole,
)
from psetw_platform.services.analytics import (
    compute_delay_risk_rows,
    compute_dependency_health,
    compute_portfolio_metrics,
    detect_activity_anomalies,
)
from psetw_platform.services.planning import (
    compute_calendar_buckets,
    compute_critical_path_codes,
    compute_phase_progress,
)

router = APIRouter(tags=["web-ui"])
templates = Jinja2Templates(
    directory=str(Path(__file__).resolve().parent.parent / "web" / "templates"),
)
SESSION_COOKIE_NAME = "psetw_ui_session"
DEFAULT_PAGE_SIZE = 50
ALLOWED_PAGE_SIZES = {25, 50, 100, 250, 1000}
LOGIN_DEMO_USERS: dict[str, tuple[str, str, str]] = {
    "planner": ("Planner", "planner", "planner123"),
    "management": ("Management", "management", "management123"),
    "technician": ("Technician", "technician", "technician123"),
}
ENGINE_UPLOAD_EXTENSIONS = {".xlsx", ".xls", ".csv", ".docx"}
ACTIVITY_EXPORT_COLUMNS: list[tuple[str, str]] = [
    ("Activity ID", "activity_code"),
    ("Activity Name", "activity_name"),
    ("Phase", "phase"),
    ("Sub Activity", "sub_activity"),
    ("Planned Start Date", "planned_start_date"),
    ("Planned End Date", "planned_end_date"),
    ("Planned Duration Hours", "planned_duration_hours"),
    ("Actual Start Date", "actual_start_date"),
    ("Actual End Date", "actual_end_date"),
    ("Actual Duration Hours", "actual_duration_hours"),
    ("Base Effort Hours", "base_effort_hours"),
    ("Required Materials", "required_materials"),
    ("Required Tools", "required_tools"),
    ("Dependencies", "dependencies"),
    ("Dependency Type", "dependency_type"),
    ("Priority", "priority"),
    ("Milestone", "milestone"),
    ("Assigned Manpower", "assigned_manpower"),
    ("Manpower Skill Level", "manpower_skill_level"),
    ("Resource Name", "resource_name"),
    ("Resource Department", "resource_department"),
    ("Shift Type", "shift_type"),
    ("Material Status", "material_status"),
    ("Material Ownership", "material_ownership"),
    ("Supplier / Vendor", "material_supplier"),
    ("Material Criticality", "material_criticality"),
    ("Material Required Date", "material_required_date"),
    ("Material Received Date", "material_received_date"),
    ("Material Lead Time", "material_lead_time"),
    ("Activity Status", "status"),
    ("Completion Percentage", "completion_percentage"),
    ("Risk Level", "risk_level"),
    ("Risk Score", "risk_score"),
    ("Risk Probability", "risk_probability"),
    ("Risk Impact", "risk_impact"),
    ("Mitigation Status", "risk_mitigation_status"),
    ("Risk Owner", "risk_owner"),
    ("Risk Review Date", "risk_review_date"),
    ("Delay Reason", "delay_reason"),
    ("Manual Override Duration", "manual_override_duration"),
    ("Override Reason", "override_reason"),
    ("Override Approved By", "override_approved_by"),
    ("Estimated Cost", "estimated_cost"),
    ("Actual Cost", "actual_cost"),
    ("Cost Center", "cost_center"),
    ("Last Modified By", "last_modified_by"),
    ("Last Modified Date", "last_modified_date"),
    ("Remarks", "remarks"),
]


def _parse_date(value: str | None) -> date | None:
    if value is None:
        return None
    raw = value.strip()
    if not raw:
        return None
    try:
        return date.fromisoformat(raw)
    except ValueError:
        try:
            return datetime.fromisoformat(raw).date()
        except ValueError:
            return None


def _serialize_activity(activity: Activity) -> dict[str, object]:
    return {
        "activity_code": activity.activity_code,
        "activity_name": activity.activity_name,
        "phase": activity.phase,
        "sub_activity": activity.sub_activity,
        "status": activity.status.value,
        "completion_percentage": activity.completion_percentage,
        "planned_start_date": activity.planned_start_date.isoformat() if activity.planned_start_date else None,
        "planned_end_date": activity.planned_end_date.isoformat() if activity.planned_end_date else None,
        "planned_duration_hours": activity.planned_duration_hours,
        "actual_start_date": activity.actual_start_date.isoformat() if activity.actual_start_date else None,
        "actual_end_date": activity.actual_end_date.isoformat() if activity.actual_end_date else None,
        "actual_duration_hours": activity.actual_duration_hours,
        "base_effort_hours": activity.base_effort_hours,
        "required_materials": activity.required_materials,
        "required_tools": activity.required_tools,
        "dependencies": activity.dependencies,
        "dependency_type": activity.dependency_type,
        "priority": activity.priority,
        "milestone": activity.milestone,
        "assigned_manpower": activity.assigned_manpower,
        "manpower_skill_level": activity.manpower_skill_level,
        "resource_name": activity.resource_name,
        "resource_department": activity.resource_department,
        "shift_type": activity.shift_type,
        "material_status": activity.material_status,
        "material_ownership": activity.material_ownership,
        "material_supplier": activity.material_supplier,
        "material_criticality": activity.material_criticality,
        "material_required_date": (
            activity.material_required_date.isoformat() if activity.material_required_date else None
        ),
        "material_received_date": (
            activity.material_received_date.isoformat() if activity.material_received_date else None
        ),
        "material_lead_time": activity.material_lead_time,
        "risk_level": activity.risk_level,
        "risk_score": activity.risk_score,
        "risk_probability": activity.risk_probability,
        "risk_impact": activity.risk_impact,
        "risk_mitigation_status": activity.risk_mitigation_status,
        "risk_owner": activity.risk_owner,
        "risk_review_date": activity.risk_review_date.isoformat() if activity.risk_review_date else None,
        "manual_override_duration": activity.manual_override_duration,
        "override_reason": activity.override_reason,
        "override_approved_by": activity.override_approved_by,
        "estimated_cost": activity.estimated_cost,
        "actual_cost": activity.actual_cost,
        "cost_center": activity.cost_center,
        "last_modified_by": activity.last_modified_by,
        "last_modified_date": activity.last_modified_date.isoformat() if activity.last_modified_date else None,
        "delay_reason": activity.delay_reason,
        "remarks": activity.remarks,
    }


def _get_cookie_user(request: Request, db: DBSession) -> User | None:
    token = request.cookies.get(SESSION_COOKIE_NAME)
    if not token:
        return None
    username = decode_access_token(token)
    if username is None:
        return None
    user = db.scalar(select(User).where(User.username == username))
    if user is None or not user.is_active:
        return None
    return user


def _login_redirect() -> RedirectResponse:
    return RedirectResponse(url="/ui/login", status_code=303)


def _ensure_projects(db: DBSession, user: User) -> list[Project]:
    projects = list(db.scalars(select(Project).order_by(Project.updated_at.desc())).all())
    if projects:
        return projects
    project = Project(name="Project 1", created_by=user.username)
    db.add(project)
    db.commit()
    db.refresh(project)
    return [project]


def _resolve_project(projects: Iterable[Project], requested_id: str | None) -> Project:
    project_list = list(projects)
    if requested_id:
        for project in project_list:
            if project.id == requested_id:
                return project
    return next(iter(project_list))


def _base_context(
    request: Request,
    user: User,
    projects: list[Project],
    active_project: Project,
    current_path: str,
    message: str = "",
) -> dict[str, object]:
    return {
        "request": request,
        "user": user,
        "projects": projects,
        "active_project": active_project,
        "current_path": current_path,
        "message": message,
        "status_values": [status.value for status in ActivityStatus],
    }


def _normalize_header(value: object) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").replace("-", " ").split())


IMPORT_HEADER_MAP = {
    "activity id": "activity_code",
    "activityid": "activity_code",
    "activity code": "activity_code",
    "activity_code": "activity_code",
    "activity name": "activity_name",
    "activity_name": "activity_name",
    "phase": "phase",
    "sub activity": "sub_activity",
    "subactivity": "sub_activity",
    "sub_activity": "sub_activity",
    "planned start date": "planned_start_date",
    "planned_start_date": "planned_start_date",
    "planned end date": "planned_end_date",
    "planned_end_date": "planned_end_date",
    "planned duration hours": "planned_duration_hours",
    "actual start date": "actual_start_date",
    "actual end date": "actual_end_date",
    "actual duration hours": "actual_duration_hours",
    "base effort hours": "base_effort_hours",
    "baseefforthours": "base_effort_hours",
    "required materials": "required_materials",
    "required tools": "required_tools",
    "dependencies": "dependencies",
    "dependency type": "dependency_type",
    "priority": "priority",
    "milestone": "milestone",
    "assigned manpower": "assigned_manpower",
    "manpower skill level": "manpower_skill_level",
    "resource name": "resource_name",
    "resource department": "resource_department",
    "shift type": "shift_type",
    "material status": "material_status",
    "material ownership": "material_ownership",
    "supplier / vendor": "material_supplier",
    "supplier": "material_supplier",
    "material criticality": "material_criticality",
    "material required date": "material_required_date",
    "material received date": "material_received_date",
    "material lead time": "material_lead_time",
    "activity status": "status",
    "status": "status",
    "completion percentage": "completion_percentage",
    "risk level": "risk_level",
    "risk score": "risk_score",
    "risk probability": "risk_probability",
    "risk impact": "risk_impact",
    "mitigation status": "risk_mitigation_status",
    "risk mitigation status": "risk_mitigation_status",
    "risk owner": "risk_owner",
    "risk review date": "risk_review_date",
    "delay reason": "delay_reason",
    "manual override duration": "manual_override_duration",
    "override reason": "override_reason",
    "override approved by": "override_approved_by",
    "estimated cost": "estimated_cost",
    "actual cost": "actual_cost",
    "cost center": "cost_center",
    "last modified by": "last_modified_by",
    "last modified date": "last_modified_date",
    "remarks": "remarks",
}

REQUIRED_IMPORT_FIELDS = {
    "activity_code",
    "activity_name",
    "phase",
    "sub_activity",
    "base_effort_hours",
    "required_materials",
    "required_tools",
    "material_ownership",
    "material_lead_time",
    "dependencies",
}

STATUS_MAP = {status.value.lower(): status for status in ActivityStatus}


def _as_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _as_int(value: object, default: int = 0) -> int:
    if value is None or value == "":
        return default
    try:
        return int(float(str(value).strip()))
    except (ValueError, TypeError):
        return default


def _as_date(value: object) -> date | None:
    if value is None:
        return None
    if isinstance(value, date):
        return value
    return _parse_date(str(value))


def _parse_dependencies(value: object) -> list[str]:
    if value is None:
        return []
    tokens = []
    for part in str(value).replace(";", ",").split(","):
        token = part.strip()
        if token:
            tokens.append(token)
    return list(dict.fromkeys(tokens))


def _coerce_status(value: object) -> ActivityStatus:
    normalized = str(value or "").strip().lower()
    return STATUS_MAP.get(normalized, ActivityStatus.not_started)


def _next_activity_code(existing_codes: set[str]) -> str:
    max_code = 0
    for code in existing_codes:
        if code.startswith("ACT-"):
            suffix = code.removeprefix("ACT-")
            if suffix.isdigit():
                max_code = max(max_code, int(suffix))
    return f"ACT-{max_code + 1:04d}"


def _activity_to_export_row(activity: Activity) -> dict[str, object]:
    row: dict[str, object] = {}
    for label, field in ACTIVITY_EXPORT_COLUMNS:
        value = getattr(activity, field)
        if isinstance(value, ActivityStatus):
            row[label] = value.value
        elif isinstance(value, list):
            row[label] = ", ".join(value)
        elif isinstance(value, date):
            row[label] = value.isoformat()
        else:
            row[label] = value
    return row


def _parse_import_rows(file_name: str, content: bytes) -> list[dict[str, object]]:
    lower_name = file_name.lower()
    if lower_name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="ignore")
        reader = csv.DictReader(io.StringIO(text))
        return [dict(row) for row in reader]
    if lower_name.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content), data_only=True)
        sheet = workbook.active
        if sheet is None:
            return []
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip() for value in rows[0]]
        parsed: list[dict[str, object]] = []
        for row in rows[1:]:
            parsed.append({headers[idx]: row[idx] if idx < len(row) else "" for idx in range(len(headers))})
        return parsed
    raise ValueError("Unsupported file format. Use .xlsx or .csv")


def _apply_import_row(activity: Activity, row_values: dict[str, object]) -> None:
    date_fields = {
        "planned_start_date",
        "planned_end_date",
        "actual_start_date",
        "actual_end_date",
        "material_required_date",
        "material_received_date",
        "risk_review_date",
        "last_modified_date",
    }
    for key, value in row_values.items():
        if key == "status":
            activity.status = _coerce_status(value)
        elif key in date_fields:
            setattr(activity, key, _as_date(value))
        elif key in {
            "completion_percentage",
            "planned_duration_hours",
            "actual_duration_hours",
            "base_effort_hours",
            "assigned_manpower",
            "material_lead_time",
            "risk_score",
            "risk_probability",
            "risk_impact",
            "manual_override_duration",
            "estimated_cost",
            "actual_cost",
        }:
            setattr(activity, key, _as_int(value))
        elif key == "dependencies":
            setattr(activity, key, _parse_dependencies(value))
        else:
            setattr(activity, key, _as_text(value))


def _compute_activity_duration_days(activity: Activity) -> int:
    if activity.planned_start_date and activity.planned_end_date:
        delta = (activity.planned_end_date - activity.planned_start_date).days
        if delta > 0:
            return delta
    if activity.planned_duration_hours > 0:
        return max(1, (activity.planned_duration_hours + 23) // 24)
    if activity.base_effort_hours > 0:
        return max(1, (activity.base_effort_hours + 23) // 24)
    return 1


def _build_calendar_matrix(
    activities: list[Activity],
    month_date: date,
    display_mode: str,
) -> list[list[dict[str, object]]]:
    month_calendar = Calendar(firstweekday=SUNDAY)
    matrix: list[list[dict[str, object]]] = []
    today = date.today()
    weeks = month_calendar.monthdatescalendar(month_date.year, month_date.month)
    if not weeks:
        return matrix

    visible_start = weeks[0][0]
    visible_end = weeks[-1][-1]
    activity_by_day: dict[date, list[Activity]] = {}
    for activity in activities:
        window = _activity_window(activity, display_mode)
        if window is None:
            continue
        start, end = window
        if end < visible_start or start > visible_end:
            continue
        current = max(start, visible_start)
        until = min(end, visible_end)
        while current <= until:
            activity_by_day.setdefault(current, []).append(activity)
            current += timedelta(days=1)

    for week in weeks:
        week_rows: list[dict[str, object]] = []
        for day in week:
            week_rows.append(
                {
                    "day": day,
                    "in_month": day.month == month_date.month,
                    "is_today": day == today,
                    "activities": sorted(
                        activity_by_day.get(day, []),
                        key=lambda activity: activity.activity_code.lower(),
                    ),
                }
            )
        matrix.append(week_rows)
    return matrix


def _extract_docx_text(content: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            xml_bytes = archive.read("word/document.xml")
    except (zipfile.BadZipFile, KeyError):
        return ""
    decoded_xml = xml_bytes.decode("utf-8", errors="ignore")
    text_parts = [match.strip() for match in re.findall(r"<w:t[^>]*>(.*?)</w:t>", decoded_xml) if match.strip()]
    return "\n".join(part for part in text_parts if part)


def _extract_xlsx_text(content: bytes) -> str:
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        return ""
    try:
        sheet = workbook.active
        if sheet is None:
            return ""
        rows: list[str] = []
        for row in sheet.iter_rows(min_row=1, max_row=40, min_col=1, max_col=12, values_only=True):
            parts = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if parts:
                rows.append(" | ".join(parts))
        return "\n".join(rows)
    finally:
        workbook.close()


def _extract_csv_text(content: bytes) -> str:
    decoded = content.decode("utf-8", errors="ignore")
    lines = [line.strip() for line in decoded.splitlines()[:200] if line.strip()]
    return "\n".join(lines)


def _extract_engine_text(file_name: str, content: bytes) -> str:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".docx":
        return _extract_docx_text(content)
    if suffix == ".xlsx":
        return _extract_xlsx_text(content)
    if suffix in {".csv", ".txt"}:
        return _extract_csv_text(content)
    return ""


def _extract_field(text: str, labels: list[str]) -> str:
    if not text:
        return ""
    label_pattern = "|".join(re.escape(label) for label in labels)
    regex = re.compile(rf"(?:{label_pattern})\s*[:\-]\s*(.+)", re.IGNORECASE)
    for line in text.splitlines():
        match = regex.search(line)
        if match:
            return match.group(1).strip()[:300]
    return ""


def _derive_engine_summary(file_name: str, extracted_text: str) -> dict[str, object]:
    stem = Path(file_name).stem.replace("_", " ").replace("-", " ").strip()
    engine_model = _extract_field(extracted_text, ["engine model", "engine name", "engine"])
    customer = _extract_field(extracted_text, ["customer", "client"])
    scope = _extract_field(extracted_text, ["scope", "program", "project"])
    remarks = _extract_field(extracted_text, ["remarks", "notes", "requirements summary"])
    required_tools = _extract_field(extracted_text, ["required tools", "tools"])
    required_materials = _extract_field(extracted_text, ["required materials", "materials"])
    preparation_phases = _extract_field(extracted_text, ["preparation phases", "phases"])
    key_milestones = _extract_field(extracted_text, ["key milestones", "milestones"])
    dependencies = _extract_field(extracted_text, ["dependencies", "depends on"])
    return {
        "engine_model": engine_model or stem,
        "customer": customer,
        "scope": scope,
        "remarks": remarks,
        "required_tools": required_tools,
        "required_materials": required_materials,
        "preparation_phases": preparation_phases,
        "key_milestones": key_milestones,
        "dependencies": dependencies,
        "preview_text": extracted_text[:3000],
    }


def _safe_ratio(numerator: int, denominator: int) -> float:
    if denominator <= 0:
        return 0.0
    return round((numerator / denominator) * 100.0, 2)


def _activity_delay_days(activity: Activity, today: date | None = None) -> int:
    anchor = today or date.today()
    if activity.planned_end_date is None:
        return 0
    if activity.status == ActivityStatus.completed:
        end_date = activity.actual_end_date or activity.planned_end_date
        return max(0, (end_date - activity.planned_end_date).days)
    return max(0, (anchor - activity.planned_end_date).days)


def _build_weekly_completion(activities: list[Activity]) -> list[dict[str, int | str]]:
    anchor = date.today()
    results: list[dict[str, int | str]] = []
    for offset in range(5, -1, -1):
        week_start = anchor - timedelta(days=anchor.weekday()) - timedelta(days=offset * 7)
        week_end = week_start + timedelta(days=6)
        completed = 0
        for activity in activities:
            completed_date = activity.actual_end_date
            if completed_date and week_start <= completed_date <= week_end:
                completed += 1
        results.append(
            {
                "label": f"{week_start.strftime('%d %b')} - {week_end.strftime('%d %b')}",
                "completed": completed,
            }
        )
    peak = max((int(row["completed"]) for row in results), default=0)
    for row in results:
        completed = int(row["completed"])
        row["bar_pct"] = int((completed / peak) * 100) if peak > 0 else 0
    return results


def _build_activities_query_params(
    project_id: str,
    search: str = "",
    status_filter: str = "",
    phase_filter: str = "",
    priority_filter: str = "",
    department_filter: str = "",
    page: int = 1,
    page_size: int = DEFAULT_PAGE_SIZE,
    column_mode: str = "core",
    highlight_id: str = "",
    message: str = "",
) -> str:
    payload: dict[str, str] = {
        "project_id": project_id,
        "search": search,
        "status_filter": status_filter,
        "phase_filter": phase_filter,
        "priority_filter": priority_filter,
        "department_filter": department_filter,
        "page": str(page),
        "page_size": str(page_size),
        "column_mode": "all" if column_mode == "all" else "core",
    }
    if highlight_id.strip():
        payload["highlight_id"] = highlight_id.strip()
    if message.strip():
        payload["message"] = message.strip()
    return urlencode(payload)


def _coerce_form_bool(value: str | None) -> bool:
    return (value or "").strip().lower() in {"1", "true", "yes", "on"}


def _normalized_move_scope(value: str) -> str:
    normalized = value.strip().lower()
    if normalized in {"actual", "both"}:
        return normalized
    return "planned"


def _status_from_raw(value: str) -> ActivityStatus:
    normalized = value.strip().lower()
    if normalized in {"on hold", "on_hold"}:
        return ActivityStatus.blocked
    return STATUS_MAP.get(normalized, ActivityStatus.not_started)


def _apply_status_business_rules(activity: Activity, next_status: ActivityStatus, working_date: date) -> list[str]:
    warnings: list[str] = []
    activity.status = next_status
    if next_status == ActivityStatus.in_progress and activity.actual_start_date is None:
        activity.actual_start_date = working_date
    if next_status == ActivityStatus.completed:
        if activity.actual_start_date is None:
            activity.actual_start_date = working_date
        if activity.actual_end_date is None:
            activity.actual_end_date = working_date
        if activity.completion_percentage < 100:
            activity.completion_percentage = 100
    if next_status == ActivityStatus.delayed and not activity.delay_reason.strip():
        warnings.append("Delayed status requires a delay reason.")
    if activity.completion_percentage == 100 and next_status != ActivityStatus.completed:
        warnings.append("Completion is 100% but status is not Completed.")
    return warnings


def _shift_activity_dates(activity: Activity, days: int, scope: str, anchor: date) -> None:
    if scope in {"planned", "both"}:
        start = activity.planned_start_date or anchor
        end = activity.planned_end_date or start
        activity.planned_start_date = start + timedelta(days=days)
        activity.planned_end_date = end + timedelta(days=days)
    if scope in {"actual", "both"}:
        start = activity.actual_start_date or activity.planned_start_date or anchor
        end = activity.actual_end_date or start
        activity.actual_start_date = start + timedelta(days=days)
        activity.actual_end_date = end + timedelta(days=days)


def _activity_window(activity: Activity, display_mode: str) -> tuple[date, date] | None:
    mode = display_mode.strip().lower()
    if mode == "actual":
        start = activity.actual_start_date or activity.actual_end_date
        end = activity.actual_end_date or activity.actual_start_date
    elif mode == "mixed":
        start = activity.actual_start_date or activity.planned_start_date or activity.planned_end_date
        end = activity.actual_end_date or activity.planned_end_date or activity.planned_start_date
    else:
        start = activity.planned_start_date or activity.planned_end_date
        end = activity.planned_end_date or activity.planned_start_date
    if start is None or end is None:
        return None
    if end < start:
        return (end, start)
    return (start, end)


def _collect_activity_warnings(
    activity: Activity,
    by_code: dict[str, Activity],
    today: date | None = None,
) -> list[str]:
    anchor = today or date.today()
    warnings: list[str] = []
    if (
        activity.planned_start_date
        and activity.planned_end_date
        and activity.planned_end_date < activity.planned_start_date
    ):
        warnings.append("Planned end date is before planned start date.")
    if (
        activity.actual_start_date
        and activity.actual_end_date
        and activity.actual_end_date < activity.actual_start_date
    ):
        warnings.append("Actual end date is before actual start date.")
    if activity.completion_percentage < 0 or activity.completion_percentage > 100:
        warnings.append("Completion percentage must be between 0 and 100.")
    if activity.status == ActivityStatus.completed and activity.completion_percentage < 100:
        warnings.append("Completed activities should usually have 100% completion.")
    if activity.status == ActivityStatus.delayed and not activity.delay_reason.strip():
        warnings.append("Delayed activity is missing delay reason.")
    if activity.status == ActivityStatus.not_started and (activity.actual_start_date or activity.actual_end_date):
        warnings.append("Not Started activity still has actual dates.")
    if activity.planned_start_date and activity.material_status.strip().lower() == "not ordered":
        if activity.planned_start_date <= anchor + timedelta(days=2):
            warnings.append("Material is not ordered while planned start is near or overdue.")
    for dependency_code in activity.dependencies:
        dependency = by_code.get(dependency_code)
        if dependency is None:
            warnings.append(f"Missing dependency: {dependency_code}.")
            continue
        if (
            dependency.planned_end_date
            and activity.planned_start_date
            and dependency.planned_end_date > activity.planned_start_date
        ):
            warnings.append(
                f"Dependency {dependency_code} ends after this activity starts."
            )
    return warnings


def _suggest_adjacent_activity_code(source_code: str, existing_codes: set[str]) -> str:
    match = re.search(r"(\d+)$", source_code)
    if not match:
        return _next_activity_code(existing_codes)
    prefix = source_code[: match.start(1)]
    value = int(match.group(1))
    width = len(match.group(1))
    for delta in range(1, 1000):
        candidate = f"{prefix}{value + delta:0{width}d}"
        if candidate not in existing_codes:
            return candidate
    return _next_activity_code(existing_codes)


@router.get("/ui", include_in_schema=False)
def ui_home(request: Request, db: DBSession) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    return RedirectResponse(url="/ui/dashboard", status_code=303)


@router.get("/ui/login", response_class=HTMLResponse, include_in_schema=False)
def ui_login_page(
    request: Request,
    demo_user: str = "",
    username: str = "",
    info: str = "",
) -> HTMLResponse:
    prefill_username = username.strip()
    prefill_credential: str | None = None
    role_hint = ""
    key = demo_user.strip().lower()
    if key in LOGIN_DEMO_USERS:
        role_hint, prefill_username, prefill_credential = LOGIN_DEMO_USERS[key]

    return templates.TemplateResponse(
        "login.html",
        {
            "request": request,
            "error": "",
            "title": "Sign In",
            "prefill_username": prefill_username,
            "prefill_credential": prefill_credential,
            "remember_me": False,
            "role_hint": role_hint,
            "info_message": info.strip(),
            "demo_cards": [
                {
                    "label": label,
                    "key": item_key,
                    "username": login_username,
                    "password": login_password,
                }
                for item_key, (label, login_username, login_password) in LOGIN_DEMO_USERS.items()
            ],
        },
    )


@router.post("/ui/login", response_class=HTMLResponse, include_in_schema=False)
def ui_login_action(
    request: Request,
    db: DBSession,
    username: Annotated[str, Form()],
    password: Annotated[str, Form()],
    remember_me: Annotated[bool, Form()] = False,
) -> Response:
    user = db.scalar(select(User).where(User.username == username.strip().lower()))
    if user is None or not verify_password(password, user.password_hash):
        return templates.TemplateResponse(
            "login.html",
            {
                "request": request,
                "error": "Invalid username or password. Check credentials and try again.",
                "title": "Sign In",
                "prefill_username": username.strip(),
                "prefill_credential": None,
                "remember_me": remember_me,
                "role_hint": "",
                "info_message": "",
                "demo_cards": [
                    {
                        "label": label,
                        "key": item_key,
                        "username": login_username,
                        "password": login_password,
                    }
                    for item_key, (label, login_username, login_password) in LOGIN_DEMO_USERS.items()
                ],
            },
            status_code=401,
        )

    expires_minutes = 60 * 24 * 30 if remember_me else 60 * 8
    token = create_access_token(subject=user.username, expires_minutes=expires_minutes)
    response = RedirectResponse(url="/ui/dashboard", status_code=303)
    response.set_cookie(
        SESSION_COOKIE_NAME,
        token,
        max_age=expires_minutes * 60,
        httponly=True,
        samesite="lax",
    )
    return response


@router.post("/ui/logout", include_in_schema=False)
def ui_logout() -> RedirectResponse:
    response = RedirectResponse(url="/ui/login", status_code=303)
    response.delete_cookie(SESSION_COOKIE_NAME)
    return response


@router.get("/ui/dashboard", response_class=HTMLResponse, include_in_schema=False)
def ui_dashboard(
    request: Request,
    db: DBSession,
    project_id: str | None = None,
    time_range: str = "30d",
) -> Response:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()

    projects = _ensure_projects(db, user)
    active_project = _resolve_project(projects, project_id)
    activities = list(
        db.scalars(
            select(Activity)
            .where(Activity.project_id == active_project.id)
            .order_by(Activity.created_at.asc())
        ).all()
    )
    actions = list(db.scalars(select(ActionItem).where(ActionItem.project_id == active_project.id)).all())

    anomalies = detect_activity_anomalies(activities)
    dependency_health = compute_dependency_health(activities)
    phase_progress = compute_phase_progress(activities)
    top_risks = compute_delay_risk_rows(activities)
    critical_path_codes = compute_critical_path_codes(activities)

    severity_counts = {"Critical": 0, "High": 0, "Medium": 0, "Low": 0}
    for anomaly in anomalies:
        severity = anomaly.severity if anomaly.severity in severity_counts else "Low"
        severity_counts[severity] += 1

    open_actions = [action for action in actions if action.status != ActionStatus.closed]
    overdue_actions = [
        action for action in open_actions if action.due_date is not None and action.due_date < date.today()
    ]
    completed_activities = [activity for activity in activities if activity.status == ActivityStatus.completed]
    estimated_cost = sum(activity.estimated_cost for activity in activities)
    actual_cost = sum(activity.actual_cost for activity in activities)
    materials_ready_count = sum(
        1
        for activity in activities
        if activity.material_status.strip().lower() in {"ready", "received", "available", "ordered"}
    )
    on_time_completed = sum(
        1
        for activity in completed_activities
        if activity.planned_end_date is None
        or (activity.actual_end_date is not None and activity.actual_end_date <= activity.planned_end_date)
    )
    overdue_milestones = sum(
        1
        for activity in activities
        if activity.milestone.strip()
        and activity.planned_end_date is not None
        and activity.planned_end_date < date.today()
        and activity.status != ActivityStatus.completed
    )
    pending_approvals = sum(
        1
        for activity in activities
        if activity.override_reason.strip() and not activity.override_approved_by.strip()
    )
    delay_by_phase: dict[str, int] = {}
    for row in top_risks:
        delay_by_phase[row.phase or "Unassigned"] = delay_by_phase.get(row.phase or "Unassigned", 0) + row.delay_days
    delay_by_phase_rows = [
        {"phase": phase, "delay_days": delay_days}
        for phase, delay_days in sorted(delay_by_phase.items(), key=lambda item: item[1], reverse=True)
    ][:8]

    material_status_rows: list[dict[str, object]] = []
    by_material_status: dict[str, int] = {}
    for activity in activities:
        label = activity.material_status.strip() or "Unknown"
        by_material_status[label] = by_material_status.get(label, 0) + 1
    for label, count in sorted(by_material_status.items(), key=lambda item: item[1], reverse=True):
        material_status_rows.append({"label": label, "count": count})

    tools_ready = sum(1 for activity in activities if activity.required_tools.strip())
    tool_readiness_pct = _safe_ratio(tools_ready, len(activities))
    upcoming_milestones = sorted(
        [
            activity
            for activity in activities
            if activity.milestone.strip()
            and activity.planned_end_date is not None
            and 0 <= (activity.planned_end_date - date.today()).days <= 14
        ],
        key=lambda activity: activity.planned_end_date or date.max,
    )
    recent_updates = sorted(
        activities,
        key=lambda activity: activity.updated_at,
        reverse=True,
    )[:8]

    context = _base_context(request, user, projects, active_project, "/ui/dashboard")
    context.update(
        {
            "title": "Executive Dashboard",
            "snapshot_date": date.today().isoformat(),
            "time_range": time_range,
            "portfolio_metrics": compute_portfolio_metrics(activities),
            "phase_progress": phase_progress,
            "top_risks": top_risks[:12],
            "dependency_health": dependency_health,
            "anomaly_counts": severity_counts,
            "open_action_count": len(open_actions),
            "overdue_action_count": len(overdue_actions),
            "estimated_cost": estimated_cost,
            "actual_cost": actual_cost,
            "cost_variance": actual_cost - estimated_cost,
            "on_time_pct": _safe_ratio(on_time_completed, len(completed_activities)),
            "materials_ready_pct": _safe_ratio(materials_ready_count, len(activities)),
            "open_anomalies": len(anomalies),
            "overdue_milestones": overdue_milestones,
            "pending_approvals": pending_approvals,
            "delay_by_phase_rows": delay_by_phase_rows,
            "material_status_rows": material_status_rows,
            "tool_readiness_pct": tool_readiness_pct,
            "weekly_completion_rows": _build_weekly_completion(activities),
            "critical_path_codes": critical_path_codes,
            "recent_updates": recent_updates,
            "upcoming_milestones": upcoming_milestones,
            "blocked_rows": [row for row in top_risks if row.blocking_dependencies][:8],
        }
    )
    return templates.TemplateResponse("dashboard.html", context)


@router.get("/ui/activities", response_class=HTMLResponse, include_in_schema=False)
def ui_activities(
    request: Request,
    db: DBSession,
    project_id: str | None = None,
    message: str = "",
    search: str = "",
    status_filter: str = "",
    phase_filter: str = "",
    priority_filter: str = "",
    department_filter: str = "",
    page: int = 1,
    page_size: int = DEFAULT_PAGE_SIZE,
    column_mode: str = "core",
    highlight_id: str = "",
) -> Response:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    projects = _ensure_projects(db, user)
    active_project = _resolve_project(projects, project_id)
    activities = list(
        db.scalars(
            select(Activity)
            .where(Activity.project_id == active_project.id)
            .order_by(Activity.created_at.asc())
        ).all()
    )
    status_options = sorted({activity.status.value for activity in activities})
    phase_options = sorted({activity.phase for activity in activities if activity.phase})
    priority_options = sorted({activity.priority for activity in activities if activity.priority.strip()})
    department_options = sorted(
        {activity.resource_department for activity in activities if activity.resource_department.strip()}
    )

    normalized_search = search.strip().lower()
    normalized_status = status_filter.strip().lower()
    normalized_phase = phase_filter.strip().lower()
    normalized_priority = priority_filter.strip().lower()
    normalized_department = department_filter.strip().lower()
    filtered = []
    for activity in activities:
        if normalized_status and activity.status.value.lower() != normalized_status:
            continue
        if normalized_phase and activity.phase.strip().lower() != normalized_phase:
            continue
        if normalized_priority and activity.priority.strip().lower() != normalized_priority:
            continue
        if normalized_department and activity.resource_department.strip().lower() != normalized_department:
            continue
        if normalized_search:
            haystack = " ".join(
                [
                    activity.activity_code,
                    activity.activity_name,
                    activity.phase,
                    activity.sub_activity,
                    activity.required_materials,
                    activity.required_tools,
                    activity.delay_reason,
                    activity.remarks,
                ]
            ).lower()
            if normalized_search not in haystack:
                continue
        filtered.append(activity)

    if page_size not in ALLOWED_PAGE_SIZES:
        page_size = DEFAULT_PAGE_SIZE
    total_count = len(activities)
    filtered_count = len(filtered)
    total_pages = max(1, (filtered_count + page_size - 1) // page_size)
    page = max(1, min(page, total_pages))
    start_idx = (page - 1) * page_size
    end_idx = start_idx + page_size
    paged_activities = filtered[start_idx:end_idx]
    by_code = {row.activity_code: row for row in activities}

    context = _base_context(request, user, projects, active_project, "/ui/activities", message)
    context.update(
        {
            "title": "Activities",
            "activities": paged_activities,
            "total_count": total_count,
            "filtered_count": filtered_count,
            "status_filter": status_filter,
            "phase_filter": phase_filter,
            "priority_filter": priority_filter,
            "department_filter": department_filter,
            "search": search,
            "status_options": status_options,
            "phase_options": phase_options,
            "priority_options": priority_options,
            "department_options": department_options,
            "page_size": page_size,
            "page": page,
            "total_pages": total_pages,
            "start_idx": start_idx,
            "end_idx": min(end_idx, filtered_count),
            "allowed_page_sizes": sorted(ALLOWED_PAGE_SIZES),
            "column_mode": "all" if column_mode == "all" else "core",
            "highlight_id": highlight_id,
            "today": date.today().isoformat(),
            "warnings_by_id": {
                activity.id: _collect_activity_warnings(
                    activity,
                    by_code,
                )
                for activity in paged_activities
            },
        }
    )
    return templates.TemplateResponse("activities.html", context)


def _activities_redirect(
    project_id: str,
    message: str = "",
    search: str = "",
    status_filter: str = "",
    phase_filter: str = "",
    priority_filter: str = "",
    department_filter: str = "",
    page: int = 1,
    page_size: int = DEFAULT_PAGE_SIZE,
    column_mode: str = "core",
    highlight_id: str = "",
) -> RedirectResponse:
    query = _build_activities_query_params(
        project_id=project_id,
        search=search,
        status_filter=status_filter,
        phase_filter=phase_filter,
        priority_filter=priority_filter,
        department_filter=department_filter,
        page=page,
        page_size=page_size,
        column_mode=column_mode,
        highlight_id=highlight_id,
        message=message,
    )
    return RedirectResponse(url=f"/ui/activities?{query}", status_code=303)


@router.post("/ui/projects/{project_id}/activities", include_in_schema=False)
def ui_create_activity(
    request: Request,
    db: DBSession,
    project_id: str,
    activity_code: Annotated[str, Form()],
    activity_name: Annotated[str, Form()],
    phase: Annotated[str, Form()] = "",
    status: Annotated[str, Form()] = ActivityStatus.not_started.value,
    planned_start_date: Annotated[str, Form()] = "",
    planned_end_date: Annotated[str, Form()] = "",
    planned_duration_hours: Annotated[int, Form()] = 0,
    completion_percentage: Annotated[int, Form()] = 0,
    sub_activity: Annotated[str, Form()] = "",
    base_effort_hours: Annotated[int, Form()] = 0,
    required_materials: Annotated[str, Form()] = "",
    required_tools: Annotated[str, Form()] = "",
    dependencies: Annotated[str, Form()] = "",
    dependency_type: Annotated[str, Form()] = "FS",
    priority: Annotated[str, Form()] = "Medium",
    milestone: Annotated[str, Form()] = "",
    assigned_manpower: Annotated[int, Form()] = 1,
    manpower_skill_level: Annotated[str, Form()] = "",
    resource_name: Annotated[str, Form()] = "",
    resource_department: Annotated[str, Form()] = "",
    shift_type: Annotated[str, Form()] = "",
    material_status: Annotated[str, Form()] = "Not Ordered",
    material_ownership: Annotated[str, Form()] = "Mechanical",
    material_supplier: Annotated[str, Form()] = "",
    material_criticality: Annotated[str, Form()] = "Medium",
    material_required_date: Annotated[str, Form()] = "",
    material_received_date: Annotated[str, Form()] = "",
    material_lead_time: Annotated[int, Form()] = 0,
    risk_level: Annotated[str, Form()] = "Low",
    risk_score: Annotated[int, Form()] = 0,
    risk_probability: Annotated[int, Form()] = 3,
    risk_impact: Annotated[int, Form()] = 3,
    risk_mitigation_status: Annotated[str, Form()] = "Planned",
    risk_owner: Annotated[str, Form()] = "",
    risk_review_date: Annotated[str, Form()] = "",
    manual_override_duration: Annotated[int, Form()] = 0,
    override_reason: Annotated[str, Form()] = "",
    override_approved_by: Annotated[str, Form()] = "",
    estimated_cost: Annotated[int, Form()] = 0,
    actual_cost: Annotated[int, Form()] = 0,
    cost_center: Annotated[str, Form()] = "",
    last_modified_by: Annotated[str, Form()] = "",
    last_modified_date: Annotated[str, Form()] = "",
    delay_reason: Annotated[str, Form()] = "",
    remarks: Annotated[str, Form()] = "",
    search: Annotated[str, Form()] = "",
    status_filter: Annotated[str, Form()] = "",
    phase_filter: Annotated[str, Form()] = "",
    priority_filter: Annotated[str, Form()] = "",
    department_filter: Annotated[str, Form()] = "",
    page: Annotated[int, Form()] = 1,
    page_size: Annotated[int, Form()] = DEFAULT_PAGE_SIZE,
    column_mode: Annotated[str, Form()] = "core",
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    if user.role not in {UserRole.planner, UserRole.management}:
        return _activities_redirect(
            project_id,
            "Only planning or management can create activities.",
            search,
            status_filter,
            phase_filter,
            priority_filter,
            department_filter,
            page,
            page_size,
            column_mode,
        )

    existing = db.scalar(
        select(Activity).where(Activity.project_id == project_id, Activity.activity_code == activity_code.strip())
    )
    if existing is not None:
        return _activities_redirect(
            project_id,
            "Activity code already exists.",
            search,
            status_filter,
            phase_filter,
            priority_filter,
            department_filter,
            page,
            page_size,
            column_mode,
        )

    try:
        status_enum = ActivityStatus(status)
    except ValueError:
        status_enum = ActivityStatus.not_started

    new_activity = Activity(
        project_id=project_id,
        activity_code=activity_code.strip(),
        activity_name=activity_name.strip() or "Unnamed",
        phase=phase.strip(),
        sub_activity=sub_activity.strip(),
        status=status_enum,
        completion_percentage=max(0, min(100, completion_percentage)),
        planned_start_date=_parse_date(planned_start_date),
        planned_end_date=_parse_date(planned_end_date),
        planned_duration_hours=max(0, planned_duration_hours),
        base_effort_hours=max(0, base_effort_hours),
        required_materials=required_materials.strip(),
        required_tools=required_tools.strip(),
        dependencies=_parse_dependencies(dependencies),
        dependency_type=dependency_type.strip() or "FS",
        priority=priority.strip() or "Medium",
        milestone=milestone.strip(),
        assigned_manpower=max(1, assigned_manpower),
        manpower_skill_level=manpower_skill_level.strip(),
        resource_name=resource_name.strip(),
        resource_department=resource_department.strip(),
        shift_type=shift_type.strip(),
        material_status=material_status.strip() or "Not Ordered",
        material_ownership=material_ownership.strip() or "Mechanical",
        material_supplier=material_supplier.strip(),
        material_criticality=material_criticality.strip() or "Medium",
        material_required_date=_parse_date(material_required_date),
        material_received_date=_parse_date(material_received_date),
        material_lead_time=max(0, material_lead_time),
        risk_level=risk_level.strip() or "Low",
        risk_score=max(0, min(100, risk_score)),
        risk_probability=max(1, min(5, risk_probability)),
        risk_impact=max(1, min(5, risk_impact)),
        risk_mitigation_status=risk_mitigation_status.strip() or "Planned",
        risk_owner=risk_owner.strip(),
        risk_review_date=_parse_date(risk_review_date),
        manual_override_duration=max(0, manual_override_duration),
        override_reason=override_reason.strip(),
        override_approved_by=override_approved_by.strip(),
        estimated_cost=max(0, estimated_cost),
        actual_cost=max(0, actual_cost),
        cost_center=cost_center.strip(),
        last_modified_by=last_modified_by.strip() or user.username,
        last_modified_date=_parse_date(last_modified_date) or date.today(),
        delay_reason=delay_reason.strip(),
        remarks=remarks.strip(),
    )
    db.add(new_activity)
    db.commit()
    return _activities_redirect(
        project_id,
        "Activity created.",
        search,
        status_filter,
        phase_filter,
        priority_filter,
        department_filter,
        page,
        page_size,
        column_mode,
        highlight_id=new_activity.id,
    )


def _insert_activity_relative(
    db: DBSession,
    project_id: str,
    anchor_activity: Activity,
    place_above: bool,
    actor: User,
) -> Activity:
    rows = list(
        db.scalars(
            select(Activity)
            .where(Activity.project_id == project_id)
            .order_by(Activity.created_at.asc(), Activity.activity_code.asc())
        ).all()
    )
    index = next((idx for idx, row in enumerate(rows) if row.id == anchor_activity.id), 0)
    if place_above:
        prev_ts = rows[index - 1].created_at if index > 0 else anchor_activity.created_at - timedelta(seconds=2)
        next_ts = anchor_activity.created_at
    else:
        prev_ts = anchor_activity.created_at
        next_ts = (
            rows[index + 1].created_at
            if index + 1 < len(rows)
            else anchor_activity.created_at + timedelta(seconds=2)
        )
    if next_ts <= prev_ts:
        next_ts = prev_ts + timedelta(seconds=2)
    midpoint = prev_ts + ((next_ts - prev_ts) / 2)
    if midpoint <= prev_ts:
        midpoint = prev_ts + timedelta(seconds=1)

    existing_codes = {row.activity_code for row in rows}
    suggested_code = _suggest_adjacent_activity_code(anchor_activity.activity_code, existing_codes)
    inserted = Activity(
        project_id=project_id,
        activity_code=suggested_code,
        activity_name="New Activity",
        phase=anchor_activity.phase,
        status=ActivityStatus.not_started,
        created_at=midpoint,
        updated_at=midpoint,
        last_modified_by=actor.username,
        last_modified_date=date.today(),
    )
    db.add(inserted)
    db.commit()
    return inserted


@router.post("/ui/projects/{project_id}/activities/{activity_id}/insert-above", include_in_schema=False)
def ui_insert_activity_above(
    request: Request,
    db: DBSession,
    project_id: str,
    activity_id: str,
    search: Annotated[str, Form()] = "",
    status_filter: Annotated[str, Form()] = "",
    phase_filter: Annotated[str, Form()] = "",
    priority_filter: Annotated[str, Form()] = "",
    department_filter: Annotated[str, Form()] = "",
    page: Annotated[int, Form()] = 1,
    page_size: Annotated[int, Form()] = DEFAULT_PAGE_SIZE,
    column_mode: Annotated[str, Form()] = "core",
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    anchor = db.scalar(select(Activity).where(Activity.project_id == project_id, Activity.id == activity_id))
    if anchor is None:
        return _activities_redirect(project_id, "Anchor row not found.")
    inserted = _insert_activity_relative(db, project_id, anchor, place_above=True, actor=user)
    return _activities_redirect(
        project_id,
        "Inserted activity above selected row.",
        search,
        status_filter,
        phase_filter,
        priority_filter,
        department_filter,
        page,
        page_size,
        column_mode,
        highlight_id=inserted.id,
    )


@router.post("/ui/projects/{project_id}/activities/{activity_id}/insert-below", include_in_schema=False)
def ui_insert_activity_below(
    request: Request,
    db: DBSession,
    project_id: str,
    activity_id: str,
    search: Annotated[str, Form()] = "",
    status_filter: Annotated[str, Form()] = "",
    phase_filter: Annotated[str, Form()] = "",
    priority_filter: Annotated[str, Form()] = "",
    department_filter: Annotated[str, Form()] = "",
    page: Annotated[int, Form()] = 1,
    page_size: Annotated[int, Form()] = DEFAULT_PAGE_SIZE,
    column_mode: Annotated[str, Form()] = "core",
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    anchor = db.scalar(select(Activity).where(Activity.project_id == project_id, Activity.id == activity_id))
    if anchor is None:
        return _activities_redirect(project_id, "Anchor row not found.")
    inserted = _insert_activity_relative(db, project_id, anchor, place_above=False, actor=user)
    return _activities_redirect(
        project_id,
        "Inserted activity below selected row.",
        search,
        status_filter,
        phase_filter,
        priority_filter,
        department_filter,
        page,
        page_size,
        column_mode,
        highlight_id=inserted.id,
    )


@router.post("/ui/projects/{project_id}/activities/{activity_id}/progress", include_in_schema=False)
def ui_update_activity_progress(
    request: Request,
    db: DBSession,
    project_id: str,
    activity_id: str,
    activity_code: Annotated[str, Form()] = "",
    activity_name: Annotated[str, Form()] = "",
    phase: Annotated[str, Form()] = "",
    sub_activity: Annotated[str, Form()] = "",
    base_effort_hours: Annotated[int, Form()] = 0,
    required_materials: Annotated[str, Form()] = "",
    required_tools: Annotated[str, Form()] = "",
    material_ownership: Annotated[str, Form()] = "",
    material_lead_time: Annotated[int, Form()] = 0,
    dependencies: Annotated[str, Form()] = "",
    planned_start_date: Annotated[str, Form()] = "",
    planned_end_date: Annotated[str, Form()] = "",
    planned_duration_hours: Annotated[int, Form()] = 0,
    priority: Annotated[str, Form()] = "",
    shift_type: Annotated[str, Form()] = "",
    material_status: Annotated[str, Form()] = "",
    material_supplier: Annotated[str, Form()] = "",
    material_required_date: Annotated[str, Form()] = "",
    material_received_date: Annotated[str, Form()] = "",
    material_criticality: Annotated[str, Form()] = "",
    actual_start_date: Annotated[str, Form()] = "",
    actual_end_date: Annotated[str, Form()] = "",
    status: Annotated[str, Form()] = ActivityStatus.not_started.value,
    quick_status: Annotated[str, Form()] = "",
    completion_percentage: Annotated[int, Form()] = 0,
    risk_level: Annotated[str, Form()] = "",
    dependency_type: Annotated[str, Form()] = "",
    override_approved_by: Annotated[str, Form()] = "",
    estimated_cost: Annotated[int, Form()] = 0,
    actual_cost: Annotated[int, Form()] = 0,
    cost_center: Annotated[str, Form()] = "",
    risk_score: Annotated[int, Form()] = 0,
    delay_reason: Annotated[str, Form()] = "",
    remarks: Annotated[str, Form()] = "",
    search: Annotated[str, Form()] = "",
    status_filter: Annotated[str, Form()] = "",
    phase_filter: Annotated[str, Form()] = "",
    priority_filter: Annotated[str, Form()] = "",
    department_filter: Annotated[str, Form()] = "",
    page: Annotated[int, Form()] = 1,
    page_size: Annotated[int, Form()] = DEFAULT_PAGE_SIZE,
    column_mode: Annotated[str, Form()] = "core",
    working_date: Annotated[str, Form()] = "",
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()

    activity = db.scalar(
        select(Activity).where(Activity.project_id == project_id, Activity.id == activity_id)
    )
    if activity is None:
        return _activities_redirect(
            project_id,
            "Activity not found.",
            search,
            status_filter,
            phase_filter,
            priority_filter,
            department_filter,
            page,
            page_size,
            column_mode,
        )

    existing_code = activity.activity_code
    proposed_code = activity_code.strip() or existing_code
    if proposed_code != existing_code:
        duplicate = db.scalar(
            select(Activity).where(
                Activity.project_id == project_id,
                Activity.activity_code == proposed_code,
                Activity.id != activity.id,
            )
        )
        if duplicate is not None:
            return _activities_redirect(
                project_id,
                "Activity ID already exists.",
                search,
                status_filter,
                phase_filter,
                priority_filter,
                department_filter,
                page,
                page_size,
                column_mode,
                highlight_id=activity.id,
            )
        activity.activity_code = proposed_code

    activity.activity_name = activity_name.strip() or activity.activity_name
    activity.phase = phase.strip()
    activity.sub_activity = sub_activity.strip()
    activity.base_effort_hours = max(0, base_effort_hours)
    activity.required_materials = required_materials.strip()
    activity.required_tools = required_tools.strip()
    activity.material_ownership = material_ownership.strip()
    activity.material_lead_time = max(0, material_lead_time)
    activity.dependencies = _parse_dependencies(dependencies)
    activity.planned_start_date = _parse_date(planned_start_date)
    activity.planned_end_date = _parse_date(planned_end_date)
    activity.planned_duration_hours = max(0, planned_duration_hours)
    activity.priority = priority.strip() or activity.priority
    activity.shift_type = shift_type.strip()
    activity.material_status = material_status.strip() or activity.material_status
    activity.material_supplier = material_supplier.strip()
    activity.material_required_date = _parse_date(material_required_date)
    activity.material_received_date = _parse_date(material_received_date)
    activity.material_criticality = material_criticality.strip() or activity.material_criticality
    activity.actual_start_date = _parse_date(actual_start_date)
    activity.actual_end_date = _parse_date(actual_end_date)

    working_anchor = _parse_date(working_date) or date.today()
    status_input = quick_status.strip() or status
    try:
        status_value = _status_from_raw(status_input)
    except ValueError:
        status_value = activity.status
    status_warnings = _apply_status_business_rules(activity, status_value, working_anchor)

    activity.completion_percentage = max(0, min(100, completion_percentage))
    activity.risk_score = max(0, min(100, risk_score))
    activity.risk_level = risk_level.strip() or activity.risk_level
    activity.dependency_type = dependency_type.strip() or activity.dependency_type
    activity.override_approved_by = override_approved_by.strip()
    activity.estimated_cost = max(0, estimated_cost)
    activity.actual_cost = max(0, actual_cost)
    activity.cost_center = cost_center.strip()
    activity.delay_reason = delay_reason.strip()
    activity.remarks = remarks.strip()

    row_warnings = _collect_activity_warnings(
        activity,
        {row.activity_code: row for row in db.scalars(select(Activity).where(Activity.project_id == project_id)).all()},
    )
    warning_message = status_warnings + row_warnings
    if warning_message:
        message = "Updated with warnings: " + " | ".join(dict.fromkeys(warning_message))
    else:
        message = "Activity updated."

    activity.last_modified_by = user.username
    activity.last_modified_date = date.today()
    db.commit()
    return _activities_redirect(
        project_id,
        message,
        search,
        status_filter,
        phase_filter,
        priority_filter,
        department_filter,
        page,
        page_size,
        column_mode,
        highlight_id=activity.id,
    )


@router.post("/ui/projects/{project_id}/activities/{activity_id}/delete", include_in_schema=False)
def ui_delete_activity(
    request: Request,
    db: DBSession,
    project_id: str,
    activity_id: str,
    search: Annotated[str, Form()] = "",
    status_filter: Annotated[str, Form()] = "",
    phase_filter: Annotated[str, Form()] = "",
    priority_filter: Annotated[str, Form()] = "",
    department_filter: Annotated[str, Form()] = "",
    page: Annotated[int, Form()] = 1,
    page_size: Annotated[int, Form()] = DEFAULT_PAGE_SIZE,
    column_mode: Annotated[str, Form()] = "core",
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    if user.role not in {UserRole.planner, UserRole.management}:
        return _activities_redirect(
            project_id,
            "Only planning or management can delete activities.",
            search,
            status_filter,
            phase_filter,
            priority_filter,
            department_filter,
            page,
            page_size,
            column_mode,
        )
    activity = db.scalar(
        select(Activity).where(Activity.project_id == project_id, Activity.id == activity_id)
    )
    if activity is not None:
        db.delete(activity)
        db.commit()
    return _activities_redirect(
        project_id,
        "Activity deleted.",
        search,
        status_filter,
        phase_filter,
        priority_filter,
        department_filter,
        page,
        page_size,
        column_mode,
    )


@router.post("/ui/projects/{project_id}/activities/bulk-status", include_in_schema=False)
def ui_bulk_status_update(
    request: Request,
    db: DBSession,
    project_id: str,
    selected_ids: Annotated[list[str], Form()],
    bulk_status: Annotated[str, Form()],
    search: Annotated[str, Form()] = "",
    status_filter: Annotated[str, Form()] = "",
    phase_filter: Annotated[str, Form()] = "",
    priority_filter: Annotated[str, Form()] = "",
    department_filter: Annotated[str, Form()] = "",
    page: Annotated[int, Form()] = 1,
    page_size: Annotated[int, Form()] = DEFAULT_PAGE_SIZE,
    column_mode: Annotated[str, Form()] = "core",
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    if not selected_ids:
        return _activities_redirect(
            project_id,
            "No rows selected.",
            search,
            status_filter,
            phase_filter,
            priority_filter,
            department_filter,
            page,
            page_size,
            column_mode,
        )
    try:
        status_value = _status_from_raw(bulk_status)
    except ValueError:
        return _activities_redirect(
            project_id,
            "Invalid status selected.",
            search,
            status_filter,
            phase_filter,
            priority_filter,
            department_filter,
            page,
            page_size,
            column_mode,
        )
    rows = list(
        db.scalars(
            select(Activity).where(Activity.project_id == project_id, Activity.id.in_(selected_ids))
        ).all()
    )
    warnings: list[str] = []
    for row in rows:
        warnings.extend(_apply_status_business_rules(row, status_value, date.today()))
        row.last_modified_by = user.username
        row.last_modified_date = date.today()
    db.commit()
    message = f"Updated status for {len(rows)} activities."
    if warnings:
        message += " Warnings: " + " | ".join(dict.fromkeys(warnings))
    return _activities_redirect(
        project_id,
        message,
        search,
        status_filter,
        phase_filter,
        priority_filter,
        department_filter,
        page,
        page_size,
        column_mode,
    )


@router.post("/ui/projects/{project_id}/activities/bulk-delete", include_in_schema=False)
def ui_bulk_delete(
    request: Request,
    db: DBSession,
    project_id: str,
    selected_ids: Annotated[list[str], Form()],
    search: Annotated[str, Form()] = "",
    status_filter: Annotated[str, Form()] = "",
    phase_filter: Annotated[str, Form()] = "",
    priority_filter: Annotated[str, Form()] = "",
    department_filter: Annotated[str, Form()] = "",
    page: Annotated[int, Form()] = 1,
    page_size: Annotated[int, Form()] = DEFAULT_PAGE_SIZE,
    column_mode: Annotated[str, Form()] = "core",
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    if user.role not in {UserRole.planner, UserRole.management}:
        return _activities_redirect(
            project_id,
            "Only planning or management can delete activities.",
            search,
            status_filter,
            phase_filter,
            priority_filter,
            department_filter,
            page,
            page_size,
            column_mode,
        )
    if not selected_ids:
        return _activities_redirect(
            project_id,
            "No rows selected.",
            search,
            status_filter,
            phase_filter,
            priority_filter,
            department_filter,
            page,
            page_size,
            column_mode,
        )
    rows = list(
        db.scalars(
            select(Activity).where(Activity.project_id == project_id, Activity.id.in_(selected_ids))
        ).all()
    )
    for row in rows:
        db.delete(row)
    db.commit()
    return _activities_redirect(
        project_id,
        f"Deleted {len(rows)} activities.",
        search,
        status_filter,
        phase_filter,
        priority_filter,
        department_filter,
        page,
        page_size,
        column_mode,
    )


@router.post("/ui/projects/{project_id}/activities/import", include_in_schema=False)
async def ui_import_activities(
    request: Request,
    db: DBSession,
    project_id: str,
    file: Annotated[UploadFile, File(...)],
    merge_strategy: Annotated[str, Form()] = "merge",
    search: Annotated[str, Form()] = "",
    status_filter: Annotated[str, Form()] = "",
    phase_filter: Annotated[str, Form()] = "",
    priority_filter: Annotated[str, Form()] = "",
    department_filter: Annotated[str, Form()] = "",
    page: Annotated[int, Form()] = 1,
    page_size: Annotated[int, Form()] = DEFAULT_PAGE_SIZE,
    column_mode: Annotated[str, Form()] = "core",
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    if user.role not in {UserRole.planner, UserRole.management}:
        return _activities_redirect(
            project_id,
            "Only planning or management can import activities.",
            search,
            status_filter,
            phase_filter,
            priority_filter,
            department_filter,
            page,
            page_size,
            column_mode,
        )
    if not file.filename:
        return _activities_redirect(
            project_id,
            "Please select a file to import.",
            search,
            status_filter,
            phase_filter,
            priority_filter,
            department_filter,
            page,
            page_size,
            column_mode,
        )

    content = await file.read()
    try:
        raw_rows = _parse_import_rows(file.filename, content)
    except ValueError as exc:
        return _activities_redirect(
            project_id,
            str(exc),
            search,
            status_filter,
            phase_filter,
            priority_filter,
            department_filter,
            page,
            page_size,
            column_mode,
        )

    if not raw_rows:
        return _activities_redirect(
            project_id,
            "No rows found in uploaded file.",
            search,
            status_filter,
            phase_filter,
            priority_filter,
            department_filter,
            page,
            page_size,
            column_mode,
        )

    normalized_rows: list[dict[str, object]] = []
    for raw_row in raw_rows:
        mapped: dict[str, object] = {}
        for header, value in raw_row.items():
            key = IMPORT_HEADER_MAP.get(_normalize_header(header))
            if key:
                mapped[key] = value
        normalized_rows.append(mapped)

    present_fields = {key for row in normalized_rows for key in row}
    missing_required = sorted(REQUIRED_IMPORT_FIELDS - present_fields)
    if missing_required:
        return _activities_redirect(
            project_id,
            "Missing required import columns: " + ", ".join(missing_required),
            search,
            status_filter,
            phase_filter,
            priority_filter,
            department_filter,
            page,
            page_size,
            column_mode,
        )

    existing = list(db.scalars(select(Activity).where(Activity.project_id == project_id)).all())
    existing_by_code = {activity.activity_code: activity for activity in existing}
    existing_codes = {activity.activity_code for activity in existing}

    if merge_strategy == "replace":
        for existing_row in existing:
            db.delete(existing_row)
        db.flush()
        existing_by_code = {}
        existing_codes = set()

    created = 0
    updated = 0
    for import_row in normalized_rows:
        code = _as_text(import_row.get("activity_code"))
        if not code:
            code = _next_activity_code(existing_codes)
        target = existing_by_code.get(code)
        if target is None:
            target = Activity(
                project_id=project_id,
                activity_code=code,
                activity_name=_as_text(import_row.get("activity_name")) or "Unnamed",
            )
            db.add(target)
            existing_by_code[code] = target
            existing_codes.add(code)
            created += 1
        else:
            updated += 1
        _apply_import_row(target, import_row)
        target.activity_code = code
        if not target.activity_name:
            target.activity_name = "Unnamed"
        target.last_modified_by = user.username
        target.last_modified_date = date.today()

    db.commit()
    return _activities_redirect(
        project_id,
        f"Import complete: {created} created, {updated} updated.",
        search,
        status_filter,
        phase_filter,
        priority_filter,
        department_filter,
        page,
        page_size,
        column_mode,
    )


def _build_export_filename(prefix: str, suffix: str) -> str:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{prefix}_{stamp}.{suffix}"


@router.get("/ui/projects/{project_id}/activities/export.csv", include_in_schema=False)
def ui_export_activities_csv(request: Request, db: DBSession, project_id: str) -> Response:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    activities = list(
        db.scalars(select(Activity).where(Activity.project_id == project_id).order_by(Activity.created_at.asc())).all()
    )
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=[label for label, _ in ACTIVITY_EXPORT_COLUMNS])
    writer.writeheader()
    for activity in activities:
        writer.writerow(_activity_to_export_row(activity))
    content = buffer.getvalue().encode("utf-8")
    headers = {"Content-Disposition": f"attachment; filename={_build_export_filename('activities', 'csv')}"}
    return Response(content=content, media_type="text/csv; charset=utf-8", headers=headers)


@router.get("/ui/projects/{project_id}/activities/export.json", include_in_schema=False)
def ui_export_activities_json(request: Request, db: DBSession, project_id: str) -> Response:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    activities = list(
        db.scalars(select(Activity).where(Activity.project_id == project_id).order_by(Activity.created_at.asc())).all()
    )
    payload = [_activity_to_export_row(activity) for activity in activities]
    headers = {"Content-Disposition": f"attachment; filename={_build_export_filename('activities', 'json')}"}
    return Response(content=json.dumps(payload, indent=2), media_type="application/json", headers=headers)


@router.get("/ui/projects/{project_id}/activities/export.xlsx", include_in_schema=False)
def ui_export_activities_xlsx(request: Request, db: DBSession, project_id: str) -> Response:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    activities = list(
        db.scalars(select(Activity).where(Activity.project_id == project_id).order_by(Activity.created_at.asc())).all()
    )

    workbook_buffer = io.BytesIO()

    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        sheet = workbook.create_sheet("Activities")
    sheet.title = "Activities"
    headers = [label for label, _ in ACTIVITY_EXPORT_COLUMNS]
    sheet.append(headers)
    for activity in activities:
        row = _activity_to_export_row(activity)
        sheet.append([row.get(header, "") for header in headers])
    workbook.save(workbook_buffer)
    workbook_buffer.seek(0)
    headers_map = {"Content-Disposition": f"attachment; filename={_build_export_filename('activities', 'xlsx')}"}
    return StreamingResponse(
        workbook_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers_map,
    )


def _render_planning_page(
    request: Request,
    db: DBSession,
    project_id: str | None,
    template_name: str,
    current_path: str,
    title: str,
    extra_context: dict[str, object],
    message: str = "",
) -> Response:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    projects = _ensure_projects(db, user)
    active_project = _resolve_project(projects, project_id)
    context = _base_context(request, user, projects, active_project, current_path, message=message)
    context.update(extra_context)
    context["title"] = title
    return templates.TemplateResponse(template_name, context)


@router.get("/ui/gantt", response_class=HTMLResponse, include_in_schema=False)
def ui_gantt(
    request: Request,
    db: DBSession,
    project_id: str | None = None,
) -> Response:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    projects = _ensure_projects(db, user)
    active_project = _resolve_project(projects, project_id)
    return RedirectResponse(
        url=f"/ui/calendar?project_id={active_project.id}&message=Gantt view was merged into interactive calendar.",
        status_code=303,
    )


@router.get("/ui/calendar", response_class=HTMLResponse, include_in_schema=False)
def ui_calendar(
    request: Request,
    db: DBSession,
    project_id: str | None = None,
    month: str | None = None,
    message: str = "",
    phase_filter: str = "",
    status_filter: str = "",
    department_filter: str = "",
    search: str = "",
    view_mode: str = "month",
    display_mode: str = "mixed",
    selected_activity_id: str = "",
    date_from: str = "",
    date_to: str = "",
) -> Response:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    projects = _ensure_projects(db, user)
    active_project = _resolve_project(projects, project_id)
    activities = list(
        db.scalars(
            select(Activity)
            .where(Activity.project_id == active_project.id)
            .order_by(Activity.created_at.asc())
        ).all()
    )
    normalized_phase = phase_filter.strip().lower()
    normalized_status = status_filter.strip().lower()
    normalized_department = department_filter.strip().lower()
    normalized_search = search.strip().lower()

    filtered: list[Activity] = []
    for activity in activities:
        if normalized_phase and activity.phase.strip().lower() != normalized_phase:
            continue
        if normalized_status and activity.status.value.lower() != normalized_status:
            continue
        if normalized_department and activity.resource_department.strip().lower() != normalized_department:
            continue
        if normalized_search:
            haystack = " ".join(
                [
                    activity.activity_code,
                    activity.activity_name,
                    activity.phase,
                    activity.sub_activity,
                    activity.resource_department,
                    activity.delay_reason,
                ]
            ).lower()
            if normalized_search not in haystack:
                continue
        filtered.append(activity)

    phase_options = sorted({activity.phase for activity in activities if activity.phase.strip()})
    status_options = [status.value for status in ActivityStatus]
    department_options = sorted(
        {
            activity.resource_department
            for activity in activities
            if activity.resource_department.strip()
        }
    )

    if month:
        try:
            month_date = datetime.strptime(f"{month}-01", "%Y-%m-%d").date()
        except ValueError:
            today = date.today()
            month_date = date(today.year, today.month, 1)
    else:
        today = date.today()
        month_date = date(today.year, today.month, 1)
    next_month = date(month_date.year + int(month_date.month == 12), (month_date.month % 12) + 1, 1)
    month_end = next_month - timedelta(days=1)
    selected_view = view_mode if view_mode in {"month", "week", "agenda"} else "month"
    selected_display = display_mode if display_mode in {"planned", "actual", "mixed"} else "mixed"

    parsed_from = _parse_date(date_from)
    parsed_to = _parse_date(date_to)
    if parsed_from and parsed_to and parsed_to < parsed_from:
        parsed_from, parsed_to = parsed_to, parsed_from
    range_start = parsed_from or month_date
    range_end = parsed_to or month_end

    calendar_rows = compute_calendar_buckets(filtered, range_start, range_end)
    calendar_matrix = _build_calendar_matrix(filtered, month_date, selected_display)

    prev_month = month_date - timedelta(days=1)
    next_month_anchor = month_end + timedelta(days=1)
    by_code = {activity.activity_code: activity for activity in activities}
    selected_activity = next((row for row in activities if row.id == selected_activity_id), None)
    selected_warnings: list[str] = []
    if selected_activity is not None:
        selected_warnings = _collect_activity_warnings(selected_activity, by_code)

    agenda_rows: list[Activity] = []
    for activity in filtered:
        window = _activity_window(activity, selected_display)
        if window is None:
            continue
        start, end = window
        if end < range_start or start > range_end:
            continue
        agenda_rows.append(activity)
    agenda_rows.sort(key=lambda row: (_activity_window(row, selected_display) or (date.max, date.max))[0])

    blocked_successors: list[Activity] = []
    if selected_activity is not None:
        for row in activities:
            if selected_activity.activity_code in row.dependencies:
                blocked_successors.append(row)

    return _render_planning_page(
        request,
        db,
        active_project.id,
        "calendar.html",
        "/ui/calendar",
        "Calendar",
        {
            "month_anchor": month_date,
            "prev_month": f"{prev_month.year}-{prev_month.month:02d}",
            "next_month": f"{next_month_anchor.year}-{next_month_anchor.month:02d}",
            "calendar_rows": calendar_rows,
            "calendar_matrix": calendar_matrix,
            "weekday_labels": ["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"],
            "selected_activity": selected_activity,
            "selected_warnings": selected_warnings,
            "blocked_successors": blocked_successors,
            "phase_filter": phase_filter,
            "status_filter": status_filter,
            "department_filter": department_filter,
            "search": search,
            "phase_options": phase_options,
            "status_options": status_options,
            "department_options": department_options,
            "view_mode": selected_view,
            "display_mode": selected_display,
            "date_from": range_start.isoformat(),
            "date_to": range_end.isoformat(),
            "agenda_rows": agenda_rows,
        },
        message=message,
    )


@router.post("/ui/projects/{project_id}/activities/{activity_id}/reschedule", include_in_schema=False)
def ui_reschedule_activity(
    request: Request,
    db: DBSession,
    project_id: str,
    activity_id: str,
    target_date: Annotated[str, Form()],
    move_scope: Annotated[str, Form()] = "planned",
    month: Annotated[str, Form()] = "",
    phase_filter: Annotated[str, Form()] = "",
    status_filter: Annotated[str, Form()] = "",
    department_filter: Annotated[str, Form()] = "",
    search: Annotated[str, Form()] = "",
    view_mode: Annotated[str, Form()] = "month",
    display_mode: Annotated[str, Form()] = "mixed",
) -> Response:
    user = _get_cookie_user(request, db)
    if user is None:
        return JSONResponse({"ok": False, "error": "Authentication required"}, status_code=401)

    target = _parse_date(target_date)
    if target is None:
        return JSONResponse({"ok": False, "error": "Invalid target date"}, status_code=422)

    activity = db.scalar(
        select(Activity).where(Activity.project_id == project_id, Activity.id == activity_id)
    )
    if activity is None:
        return JSONResponse({"ok": False, "error": "Activity not found"}, status_code=404)

    scope = _normalized_move_scope(move_scope)
    duration_days = _compute_activity_duration_days(activity)
    if scope in {"planned", "both"}:
        activity.planned_start_date = target
        activity.planned_end_date = target + timedelta(days=duration_days)
    if scope in {"actual", "both"}:
        actual_duration = max(0, (activity.actual_end_date - activity.actual_start_date).days) if (
            activity.actual_start_date and activity.actual_end_date
        ) else duration_days
        activity.actual_start_date = target
        activity.actual_end_date = target + timedelta(days=actual_duration)
        if activity.status == ActivityStatus.not_started:
            activity.status = ActivityStatus.in_progress
    activity.last_modified_by = user.username
    activity.last_modified_date = date.today()
    db.commit()
    if request.headers.get("x-requested-with", "").lower() != "xmlhttprequest":
        month_param = month.strip() or f"{target.year}-{target.month:02d}"
        return RedirectResponse(
            url=(
                "/ui/calendar?"
                + urlencode(
                    {
                        "project_id": project_id,
                        "month": month_param,
                        "message": "Activity rescheduled.",
                        "selected_activity_id": activity_id,
                        "phase_filter": phase_filter,
                        "status_filter": status_filter,
                        "department_filter": department_filter,
                        "search": search,
                        "view_mode": view_mode,
                        "display_mode": display_mode,
                    }
                )
            ),
            status_code=303,
        )
    payload = {
        "ok": True,
        "start": (activity.planned_start_date or target).isoformat(),
        "end": (activity.planned_end_date or target).isoformat(),
    }
    return JSONResponse(payload)


def _calendar_redirect(
    project_id: str,
    month: str,
    message: str,
    phase_filter: str,
    status_filter: str,
    department_filter: str,
    search: str,
    view_mode: str,
    display_mode: str,
    selected_activity_id: str = "",
    date_from: str = "",
    date_to: str = "",
) -> RedirectResponse:
    query = urlencode(
        {
            "project_id": project_id,
            "month": month,
            "message": message,
            "phase_filter": phase_filter,
            "status_filter": status_filter,
            "department_filter": department_filter,
            "search": search,
            "view_mode": view_mode,
            "display_mode": display_mode,
            "selected_activity_id": selected_activity_id,
            "date_from": date_from,
            "date_to": date_to,
        }
    )
    return RedirectResponse(url=f"/ui/calendar?{query}", status_code=303)


@router.post("/ui/projects/{project_id}/activities/{activity_id}/calendar-shift", include_in_schema=False)
def ui_calendar_shift_activity(
    request: Request,
    db: DBSession,
    project_id: str,
    activity_id: str,
    shift_days: Annotated[int, Form()] = 0,
    move_scope: Annotated[str, Form()] = "planned",
    set_start_date: Annotated[str, Form()] = "",
    set_end_date: Annotated[str, Form()] = "",
    set_actual_start_date: Annotated[str, Form()] = "",
    set_actual_end_date: Annotated[str, Form()] = "",
    mark_started: Annotated[str, Form()] = "",
    mark_completed: Annotated[str, Form()] = "",
    month: Annotated[str, Form()] = "",
    phase_filter: Annotated[str, Form()] = "",
    status_filter: Annotated[str, Form()] = "",
    department_filter: Annotated[str, Form()] = "",
    search: Annotated[str, Form()] = "",
    view_mode: Annotated[str, Form()] = "month",
    display_mode: Annotated[str, Form()] = "mixed",
    date_from: Annotated[str, Form()] = "",
    date_to: Annotated[str, Form()] = "",
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    activity = db.scalar(select(Activity).where(Activity.project_id == project_id, Activity.id == activity_id))
    if activity is None:
        return _calendar_redirect(
            project_id,
            month,
            "Activity not found.",
            phase_filter,
            status_filter,
            department_filter,
            search,
            view_mode,
            display_mode,
            activity_id,
            date_from,
            date_to,
        )

    scope = _normalized_move_scope(move_scope)
    anchor = date.today()
    if shift_days != 0:
        _shift_activity_dates(activity, shift_days, scope, anchor)

    planned_start = _parse_date(set_start_date)
    planned_end = _parse_date(set_end_date)
    actual_start = _parse_date(set_actual_start_date)
    actual_end = _parse_date(set_actual_end_date)
    if planned_start:
        activity.planned_start_date = planned_start
    if planned_end:
        activity.planned_end_date = planned_end
    if actual_start:
        activity.actual_start_date = actual_start
    if actual_end:
        activity.actual_end_date = actual_end
    if _coerce_form_bool(mark_started):
        _apply_status_business_rules(activity, ActivityStatus.in_progress, anchor)
    if _coerce_form_bool(mark_completed):
        _apply_status_business_rules(activity, ActivityStatus.completed, anchor)

    warnings = _collect_activity_warnings(
        activity,
        {row.activity_code: row for row in db.scalars(select(Activity).where(Activity.project_id == project_id)).all()},
    )
    activity.last_modified_by = user.username
    activity.last_modified_date = date.today()
    db.commit()

    message = "Calendar update applied."
    if warnings:
        message += " Warnings: " + " | ".join(dict.fromkeys(warnings))
    return _calendar_redirect(
        project_id,
        month,
        message,
        phase_filter,
        status_filter,
        department_filter,
        search,
        view_mode,
        display_mode,
        activity_id,
        date_from,
        date_to,
    )


def _delay_optimization_redirect(project_id: str, message: str = "") -> RedirectResponse:
    url = f"/ui/delay-optimization?project_id={project_id}"
    if message:
        url += f"&message={message}"
    return RedirectResponse(url=url, status_code=303)


@router.get("/ui/delay-optimization", response_class=HTMLResponse, include_in_schema=False)
def ui_delay_optimization(
    request: Request,
    db: DBSession,
    project_id: str | None = None,
    message: str = "",
) -> Response:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    projects = _ensure_projects(db, user)
    active_project = _resolve_project(projects, project_id)
    activities = list(db.scalars(select(Activity).where(Activity.project_id == active_project.id)).all())
    risk_rows = compute_delay_risk_rows(activities)
    dependency_health = compute_dependency_health(activities)
    delayed_rows = [row for row in risk_rows if row.status == ActivityStatus.delayed]
    risk_score_rows = sorted(delayed_rows, key=lambda row: row.risk_score, reverse=True)[:12]
    trend_rows: list[dict[str, object]] = []
    by_phase: dict[str, dict[str, int]] = {}
    for row in delayed_rows:
        phase = row.phase or "Unassigned"
        if phase not in by_phase:
            by_phase[phase] = {"count": 0, "delay_days": 0}
        by_phase[phase]["count"] += 1
        by_phase[phase]["delay_days"] += row.delay_days
    for phase, values in sorted(by_phase.items(), key=lambda item: item[1]["delay_days"], reverse=True):
        trend_rows.append(
            {
                "phase": phase,
                "activity_count": values["count"],
                "delay_days": values["delay_days"],
                "avg_delay": round(values["delay_days"] / values["count"], 2) if values["count"] else 0.0,
            }
        )
    return _render_planning_page(
        request,
        db,
        active_project.id,
        "delay_optimization.html",
        "/ui/delay-optimization",
        "Delay & Optimization",
        {
            "risk_rows": risk_rows,
            "delayed_rows": delayed_rows,
            "risk_score_rows": risk_score_rows,
            "dependency_health": dependency_health,
            "trend_rows": trend_rows,
        },
        message=message,
    )


@router.get("/ui/risk-register", response_class=HTMLResponse, include_in_schema=False)
def ui_risk_register_legacy(
    request: Request,
    db: DBSession,
    project_id: str | None = None,
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    projects = _ensure_projects(db, user)
    active_project = _resolve_project(projects, project_id)
    return RedirectResponse(url=f"/ui/delay-optimization?project_id={active_project.id}", status_code=303)


@router.post("/ui/projects/{project_id}/delay-root-cause", include_in_schema=False)
def ui_update_delay_root_cause(
    request: Request,
    db: DBSession,
    project_id: str,
    activity_id: Annotated[str, Form()],
    status: Annotated[str, Form()] = "",
    completion_percentage: Annotated[int, Form()] = 0,
    delay_reason: Annotated[str, Form()] = "",
    remarks: Annotated[str, Form()] = "",
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    activity = db.scalar(select(Activity).where(Activity.project_id == project_id, Activity.id == activity_id))
    if activity is None:
        return _delay_optimization_redirect(project_id, "Selected activity not found.")
    if activity.status != ActivityStatus.delayed:
        return _delay_optimization_redirect(
            project_id,
            "Only activities currently in Delayed status can be updated from this panel.",
        )
    if status:
        try:
            activity.status = _status_from_raw(status)
        except ValueError:
            pass
    activity.completion_percentage = max(0, min(100, completion_percentage))
    activity.delay_reason = delay_reason.strip()
    activity.remarks = remarks.strip()
    activity.last_modified_by = user.username
    activity.last_modified_date = date.today()
    db.commit()
    return _delay_optimization_redirect(project_id, "Delay root cause updated.")


def _engine_redirect(project_id: str, message: str = "") -> RedirectResponse:
    url = f"/ui/engine-description?project_id={project_id}"
    if message:
        url += f"&message={message}"
    return RedirectResponse(url=url, status_code=303)


@router.get("/ui/engine-description", response_class=HTMLResponse, include_in_schema=False)
def ui_engine_description(
    request: Request,
    db: DBSession,
    project_id: str | None = None,
    message: str = "",
    search: str = "",
    customer_filter: str = "",
) -> Response:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    projects = _ensure_projects(db, user)
    active_project = _resolve_project(projects, project_id)

    documents = list(
        db.scalars(
            select(EngineDocument)
            .where(EngineDocument.project_id == active_project.id)
            .order_by(EngineDocument.updated_at.desc())
        ).all()
    )

    normalized_search = search.strip().lower()
    normalized_customer = customer_filter.strip().lower()
    filtered_documents: list[EngineDocument] = []
    for document in documents:
        if normalized_customer and document.customer.strip().lower() != normalized_customer:
            continue
        if normalized_search:
            haystack = " ".join(
                [
                    document.filename,
                    document.engine_model,
                    document.customer,
                    document.scope,
                    document.remarks,
                ]
            ).lower()
            if normalized_search not in haystack:
                continue
        filtered_documents.append(document)

    customer_options = sorted({document.customer for document in documents if document.customer.strip()})
    return _render_planning_page(
        request,
        db,
        active_project.id,
        "engine_description.html",
        "/ui/engine-description",
        "Engine Description",
        {
            "documents": filtered_documents,
            "search": search,
            "customer_filter": customer_filter,
            "customer_options": customer_options,
        },
        message=message,
    )


@router.post("/ui/projects/{project_id}/engine-documents/upload", include_in_schema=False)
async def ui_upload_engine_document(
    request: Request,
    db: DBSession,
    project_id: str,
    file: Annotated[UploadFile, File(...)],
    engine_model: Annotated[str, Form()] = "",
    customer: Annotated[str, Form()] = "",
    scope: Annotated[str, Form()] = "",
    remarks: Annotated[str, Form()] = "",
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    if not file.filename:
        return _engine_redirect(project_id, "Select a document to upload.")

    suffix = Path(file.filename).suffix.lower()
    if suffix not in ENGINE_UPLOAD_EXTENSIONS:
        return _engine_redirect(project_id, "Allowed formats: .xlsx, .xls, .csv, .docx")

    content = await file.read()
    if not content:
        return _engine_redirect(project_id, "Uploaded file is empty.")

    extracted_text = _extract_engine_text(file.filename, content)
    parsed_summary = _derive_engine_summary(file.filename, extracted_text)
    document = EngineDocument(
        project_id=project_id,
        filename=Path(file.filename).name,
        content_type=file.content_type or "application/octet-stream",
        content_blob=content,
        engine_model=engine_model.strip() or _as_text(parsed_summary.get("engine_model")),
        customer=customer.strip() or _as_text(parsed_summary.get("customer")),
        scope=scope.strip() or _as_text(parsed_summary.get("scope")),
        remarks=remarks.strip() or _as_text(parsed_summary.get("remarks")),
        parsed_summary=parsed_summary,
        created_by=user.username,
    )
    db.add(document)
    db.commit()
    return _engine_redirect(project_id, "Engine requirement document uploaded.")


@router.post("/ui/projects/{project_id}/engine-documents/{document_id}/summary", include_in_schema=False)
def ui_update_engine_document_summary(
    request: Request,
    db: DBSession,
    project_id: str,
    document_id: str,
    engine_model: Annotated[str, Form()] = "",
    customer: Annotated[str, Form()] = "",
    scope: Annotated[str, Form()] = "",
    remarks: Annotated[str, Form()] = "",
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    document = db.scalar(
        select(EngineDocument).where(EngineDocument.project_id == project_id, EngineDocument.id == document_id)
    )
    if document is None:
        return _engine_redirect(project_id, "Document not found.")
    document.engine_model = engine_model.strip()
    document.customer = customer.strip()
    document.scope = scope.strip()
    document.remarks = remarks.strip()
    summary = dict(document.parsed_summary)
    summary["engine_model"] = document.engine_model
    summary["customer"] = document.customer
    summary["scope"] = document.scope
    summary["remarks"] = document.remarks
    document.parsed_summary = summary
    db.commit()
    return _engine_redirect(project_id, "Summary fields updated.")


@router.get("/ui/projects/{project_id}/engine-documents/{document_id}/download", include_in_schema=False)
def ui_download_engine_document(
    request: Request,
    db: DBSession,
    project_id: str,
    document_id: str,
) -> Response:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    document = db.scalar(
        select(EngineDocument).where(EngineDocument.project_id == project_id, EngineDocument.id == document_id)
    )
    if document is None:
        return _engine_redirect(project_id, "Document not found.")
    headers = {"Content-Disposition": f'attachment; filename="{document.filename}"'}
    return Response(content=document.content_blob, media_type=document.content_type, headers=headers)


@router.get("/ui/settings", response_class=HTMLResponse, include_in_schema=False)
def ui_settings(
    request: Request,
    db: DBSession,
    project_id: str | None = None,
) -> Response:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    projects = _ensure_projects(db, user)
    active_project = _resolve_project(projects, project_id)
    return _render_planning_page(
        request,
        db,
        active_project.id,
        "settings.html",
        "/ui/settings",
        "Settings",
        {"today": date.today().isoformat()},
    )


@router.get("/ui/anomaly-center", response_class=HTMLResponse, include_in_schema=False)
def ui_anomaly_center(
    request: Request,
    db: DBSession,
    project_id: str | None = None,
) -> Response:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    projects = _ensure_projects(db, user)
    active_project = _resolve_project(projects, project_id)
    activities = list(db.scalars(select(Activity).where(Activity.project_id == active_project.id)).all())
    anomalies = detect_activity_anomalies(activities)
    baselines = list(
        db.scalars(
            select(Baseline)
            .where(Baseline.project_id == active_project.id)
            .order_by(Baseline.created_at.desc())
        ).all()
    )
    actions = list(
        db.scalars(
            select(ActionItem)
            .where(ActionItem.project_id == active_project.id)
            .order_by(ActionItem.updated_at.desc())
        ).all()
    )
    return _render_planning_page(
        request,
        db,
        active_project.id,
        "anomaly_center.html",
        "/ui/anomaly-center",
        "Anomaly / Actions",
        {
            "anomalies": anomalies,
            "baselines": baselines,
            "actions": actions,
            "action_status_values": [status.value for status in ActionStatus],
            "action_priority_values": [priority.value for priority in ActionPriority],
            "activities": activities,
        },
    )


def _anomaly_redirect(project_id: str, message: str = "") -> RedirectResponse:
    url = f"/ui/anomaly-center?project_id={project_id}"
    if message:
        url += f"&message={message}"
    return RedirectResponse(url=url, status_code=303)


def _eod_redirect(project_id: str, message: str = "") -> RedirectResponse:
    url = f"/ui/eod-logs?project_id={project_id}"
    if message:
        url += f"&message={message}"
    return RedirectResponse(url=url, status_code=303)


@router.post("/ui/projects/{project_id}/baselines", include_in_schema=False)
def ui_create_baseline(
    request: Request,
    db: DBSession,
    project_id: str,
    name: Annotated[str, Form()] = "",
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    if user.role not in {UserRole.planner, UserRole.management}:
        return _anomaly_redirect(project_id, "Only planning or management can lock baselines.")
    activities = list(db.scalars(select(Activity).where(Activity.project_id == project_id)).all())
    baseline = Baseline(
        project_id=project_id,
        name=name.strip() or f"Baseline {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        snapshot={"activities": [_serialize_activity(activity) for activity in activities]},
        created_by=user.username,
    )
    db.add(baseline)
    db.commit()
    return _anomaly_redirect(project_id, "Baseline locked.")


@router.post("/ui/projects/{project_id}/baselines/{baseline_id}/restore", include_in_schema=False)
def ui_restore_baseline(request: Request, db: DBSession, project_id: str, baseline_id: str) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    if user.role not in {UserRole.planner, UserRole.management}:
        return _anomaly_redirect(project_id, "Only planning or management can restore baselines.")

    baseline = db.scalar(select(Baseline).where(Baseline.project_id == project_id, Baseline.id == baseline_id))
    if baseline is None:
        return _anomaly_redirect(project_id, "Baseline not found.")

    snapshot_rows = baseline.snapshot.get("activities")
    if not isinstance(snapshot_rows, list):
        return _anomaly_redirect(project_id, "Baseline snapshot is invalid.")

    current_rows = list(db.scalars(select(Activity).where(Activity.project_id == project_id)).all())
    for row in current_rows:
        db.delete(row)
    db.flush()
    existing_codes: set[str] = set()
    for snapshot_row in snapshot_rows:
        if not isinstance(snapshot_row, dict):
            continue
        code = _as_text(snapshot_row.get("activity_code")) or _next_activity_code(existing_codes)
        activity = Activity(
            project_id=project_id,
            activity_code=code,
            activity_name=_as_text(snapshot_row.get("activity_name")) or "Unnamed",
        )
        _apply_import_row(activity, snapshot_row)
        existing_codes.add(code)
        db.add(activity)
    db.commit()
    return _anomaly_redirect(project_id, f"Restored baseline {baseline.name}.")


@router.post("/ui/projects/{project_id}/actions", include_in_schema=False)
def ui_create_action(
    request: Request,
    db: DBSession,
    project_id: str,
    activity_id: Annotated[str, Form()] = "",
    title: Annotated[str, Form()] = "",
    owner: Annotated[str, Form()] = "",
    due_date: Annotated[str, Form()] = "",
    priority: Annotated[str, Form()] = "Medium",
    status: Annotated[str, Form()] = "Open",
    notes: Annotated[str, Form()] = "",
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    if user.role not in {UserRole.planner, UserRole.management}:
        return _anomaly_redirect(project_id, "Only planning or management can create actions.")
    if not title.strip() or not owner.strip():
        return _anomaly_redirect(project_id, "Action title and owner are required.")

    try:
        priority_value = ActionPriority(priority)
    except ValueError:
        priority_value = ActionPriority.medium
    try:
        status_value = ActionStatus(status)
    except ValueError:
        status_value = ActionStatus.open

    linked_activity_id = activity_id.strip() or None
    if linked_activity_id:
        exists = db.scalar(
            select(Activity.id).where(Activity.project_id == project_id, Activity.id == linked_activity_id)
        )
        if exists is None:
            return _anomaly_redirect(project_id, "Linked activity not found.")

    action = ActionItem(
        project_id=project_id,
        activity_id=linked_activity_id,
        title=title.strip(),
        owner=owner.strip(),
        due_date=_parse_date(due_date),
        priority=priority_value,
        status=status_value,
        notes=notes.strip(),
        created_by=user.username,
    )
    db.add(action)
    db.commit()
    return _anomaly_redirect(project_id, "Action created.")


@router.post("/ui/projects/{project_id}/actions/{action_id}/status", include_in_schema=False)
def ui_update_action_status(
    request: Request,
    db: DBSession,
    project_id: str,
    action_id: str,
    status: Annotated[str, Form()],
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    action = db.scalar(select(ActionItem).where(ActionItem.project_id == project_id, ActionItem.id == action_id))
    if action is None:
        return _anomaly_redirect(project_id, "Action not found.")
    try:
        action.status = ActionStatus(status)
    except ValueError:
        return _anomaly_redirect(project_id, "Invalid action status.")
    db.commit()
    return _anomaly_redirect(project_id, "Action status updated.")


@router.post("/ui/projects/{project_id}/actions/{action_id}/delete", include_in_schema=False)
def ui_delete_action(request: Request, db: DBSession, project_id: str, action_id: str) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    if user.role not in {UserRole.planner, UserRole.management}:
        return _anomaly_redirect(project_id, "Only planning or management can delete actions.")
    action = db.scalar(select(ActionItem).where(ActionItem.project_id == project_id, ActionItem.id == action_id))
    if action is not None:
        db.delete(action)
        db.commit()
    return _anomaly_redirect(project_id, "Action deleted.")


@router.get("/ui/eod-logs", response_class=HTMLResponse, include_in_schema=False)
def ui_eod_logs(
    request: Request,
    db: DBSession,
    project_id: str | None = None,
    message: str = "",
) -> Response:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    projects = _ensure_projects(db, user)
    active_project = _resolve_project(projects, project_id)
    logs = list(
        db.scalars(
            select(EodLog)
            .where(EodLog.project_id == active_project.id)
            .order_by(EodLog.log_date.desc(), EodLog.created_at.desc())
        ).all()
    )
    activities = list(
        db.scalars(
            select(Activity)
            .where(Activity.project_id == active_project.id)
            .order_by(Activity.activity_code.asc())
        ).all()
    )
    return _render_planning_page(
        request,
        db,
        active_project.id,
        "eod_logs.html",
        "/ui/eod-logs",
        "EOD Logs",
        {"logs": logs, "activities": activities},
        message=message,
    )


@router.post("/ui/projects/{project_id}/eod-logs", include_in_schema=False)
def ui_create_eod_log(
    request: Request,
    db: DBSession,
    project_id: str,
    log_date: Annotated[str, Form()],
    engineer: Annotated[str, Form()] = "",
    activities_worked_on: Annotated[str, Form()] = "",
    phase: Annotated[str, Form()] = "",
    activity_count: Annotated[int, Form()] = 0,
    hours_logged: Annotated[int, Form()] = 0,
    progress_delta: Annotated[int, Form()] = 0,
    blockers: Annotated[str, Form()] = "",
    next_day_plan: Annotated[str, Form()] = "",
    materials_received: Annotated[str, Form()] = "",
    issues_observed: Annotated[str, Form()] = "",
    status: Annotated[str, Form()] = "Open",
    verified_by: Annotated[str, Form()] = "",
) -> RedirectResponse:
    user = _get_cookie_user(request, db)
    if user is None:
        return _login_redirect()
    parsed_log_date = _parse_date(log_date) or date.today()
    entry = EodLog(
        project_id=project_id,
        log_date=parsed_log_date,
        engineer=engineer.strip() or user.display_name or user.username,
        activities_worked_on=activities_worked_on.strip(),
        phase=phase.strip(),
        activity_count=max(0, activity_count),
        hours_logged=max(0, hours_logged),
        progress_delta=progress_delta,
        blockers=blockers.strip(),
        next_day_plan=next_day_plan.strip(),
        materials_received=materials_received.strip(),
        issues_observed=issues_observed.strip(),
        status=status.strip() or "Open",
        verified_by=verified_by.strip(),
    )
    db.add(entry)
    db.commit()
    return _eod_redirect(project_id, "EOD log entry added.")
